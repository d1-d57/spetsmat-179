"""The whole conduit as a workbook: one sheet per листок, students down, problems across.

This is what makes the bot's data leaveable.  The owner ran someone else's bot in 2023 and
could not get his own tables back out of it; a school year of check-offs that exists only
inside one program is a year that can be lost by that program.  The export is therefore not
a convenience feature -- it is the exit.

THREE THINGS THIS FILE REFUSES TO DO, each of them a way an export goes wrong:

(a) IT DOES NOT RECOMPUTE THE PROJECTION.  The state of a cell is the last event for the
    pair, and that rule already lives in ``core/services/progress.py``.  A second, parallel
    reading of the journal here would be a second opinion about what a plus means, and the
    day the two disagree the teacher believes the screen while the owner believes the
    workbook.  So the export goes through ``ProgressService.states_for_many`` -- the same
    call the grids on the screen go through.

(b) IT DOES NOT INVENT SIGNS.  The alphabet is the conduit's own, copied from
    ``VALUE_REGISTRY`` in ``tools/import_konduit.py`` rather than chosen here: ``1`` is
    credited, ``x`` is handed in and not credited, an empty cell is nothing written.  The
    person who read the paper tables reads this workbook without being taught anything, and
    the importer can read it back.

(c) IT DOES NOT OPEN THE DATABASE FOR WRITING.  Teachers mark during the lesson, and the
    export is a READ.  Read-only is not a promise in a comment here, it is the connection
    URI ``file:...?mode=ro``: a stray ``insert`` raises "attempt to write a readonly
    database" instead of landing in the file somebody is tapping into.  No ``journal_mode``
    pragma is set either -- setting one is itself a write.

THE WORKBOOK IS NEVER COMMITTED.  It carries the surnames of fifty-six children.  It is
written under ``config.DB_PATH.parent`` (``data/``), which ``.gitignore`` already excludes,
and the export prints the path it wrote so that the file is found by the person who asked
for it rather than by ``git status``.

    python3 tools/export_xlsx.py            # the live database
    python3 tools/export_xlsx.py --proba    # a probe database built from migrations + seed
"""

from __future__ import annotations

import argparse
import sqlite3
import sys
import tempfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config
from core.models import CellState
from core.services.progress import ProgressService
from core.services.seeding import seed_catalogue
from infra.db import apply_migrations, connect
from infra.repositories import SqliteCatalogue, SqliteMarkJournal

#: Where the workbook is written.  Derived from ``config.DB_PATH`` and not spelled out as a
#: literal, because ``data/`` is the directory ``.gitignore`` already excludes and the
#: export must land somewhere git cannot pick it up by accident.
#:
#: 🔴 THIS BELONGS IN ``config.py`` as ``EXPORT_DIR``.  It is here because the зона of this
#: position forbids editing ``config.py``, and the debt is named in ``## ОТЧЁТ`` rather than
#: worked around by scattering a path literal through the file.
EXPORT_DIR = config.DB_PATH.parent

#: The conduit's own alphabet.  NOT invented here -- it is the reading side of
#: ``VALUE_REGISTRY`` in ``tools/import_konduit.py``: there ``1`` means solved and ``x``
#: means withdrawn, and here the same two characters mean the same two things.  A RETRACTED
#: cell is "handed in, not credited, and not a debt", which is exactly what ``x`` meant on
#: paper.
SIGN = {
    CellState.SOLVED: "1",
    CellState.RETRACTED: "x",
    CellState.EMPTY: "",
}

#: The header of the first column.  One column, not two: the paper tables carried the name
#: as one string and splitting it here would make the workbook a different shape from the
#: thing it is a copy of.
NAME_HEADER = "Ученик"


class DatabaseMissing(Exception):
    """The database to export does not exist.  Said in words, not as a traceback."""


@dataclass(frozen=True)
class ExportCounts:
    """What one export run actually wrote, so the caller prints measured numbers.

    ``cells`` is students × problems over every sheet -- every cell the workbook contains,
    including the empty ones.  Counting only the filled cells would make a half-written
    export indistinguishable from a quiet year.
    """

    sheets: int = 0
    students: int = 0
    cells: int = 0

    def __str__(self) -> str:
        return "листов %d, учеников %d, клеток %d" % (self.sheets, self.students, self.cells)


# ------------------------------------------------------------------------- the database


def open_read_only(db_path: Optional[Path] = None) -> sqlite3.Connection:
    """Open the database for reading and for nothing else.

    ``mode=ro`` is the carrier: SQLite itself refuses the write, so "the export is safe to
    run while teachers are marking" is enforced by the connection rather than by everyone
    downstream remembering.  ``busy_timeout`` is still set -- it is a reader's pragma and a
    reader that gives up instantly on a busy database is a reader that fails during exactly
    the lesson it was supposed to be safe in.
    """
    path = Path(db_path) if db_path is not None else config.DB_PATH
    if not path.exists():
        raise DatabaseMissing(
            "нет базы: %s\n"
            "живая база лежит вне git (`data/` в .gitignore), и в свежей рабочей папке её "
            "нет по построению.\n"
            "для проверки самого экспорта: python3 tools/export_xlsx.py --proba" % path
        )
    connection = sqlite3.connect("file:%s?mode=ro" % path, uri=True, isolation_level=None)
    connection.row_factory = sqlite3.Row
    connection.execute("pragma busy_timeout = %d" % config.BUSY_TIMEOUT_MS)
    return connection


def build_probe_database(directory: Path) -> Path:
    """A migrated, seeded database in ``directory`` -- the catalogue and nobody's marks.

    ``--proba`` needs something to export, and the live database is absent from every
    checkout by construction.  Building one from the migrations and the seed exercises the
    real path end to end (real schema, real 18 sheets, real 56 students) and invents no
    marks: an export of a catalogue with an empty journal is a workbook of empty cells,
    which is a true statement about a database that has none.
    """
    path = directory / "proba.db"
    apply_migrations(path, config.MIGRATIONS_DIR)
    connection = connect(path)
    try:
        seed_catalogue(connection)
    finally:
        connection.close()
    return path


# --------------------------------------------------------------------------- the export


def export(connection: sqlite3.Connection, out_path: Path) -> ExportCounts:
    """Write the whole conduit to ``out_path`` and report what was written.

    One worksheet per листок, in ``ord`` order; students down the rows in the catalogue's
    own order (surname, name); problems across the columns in ``ord`` order.  Sheets with
    no problems still get a worksheet: a листок that is missing from the workbook looks
    like a листок that never existed.
    """
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    catalogue = SqliteCatalogue(connection)
    journal = SqliteMarkJournal(connection)
    progress = ProgressService(journal, catalogue)

    students = catalogue.students()
    student_ids = [student.id for student in students]
    sheets = catalogue.sheets()

    workbook = Workbook()
    workbook.remove(workbook.active)

    cells = 0
    for sheet in sheets:
        problems = catalogue.problems_of_sheet(sheet.id)
        worksheet = workbook.create_sheet(title=_tab_title(sheet.number))

        worksheet.cell(row=1, column=1, value=NAME_HEADER)
        for column, problem in enumerate(problems, start=2):
            worksheet.cell(row=1, column=column, value=problem.label)

        states = progress.states_for_many(student_ids, [p.id for p in problems])
        for row, student in enumerate(students, start=2):
            worksheet.cell(row=row, column=1, value="%s %s" % (student.surname, student.name))
            for column, problem in enumerate(problems, start=2):
                sign = SIGN[states[(student.id, problem.id)]]
                if sign:
                    worksheet.cell(row=row, column=column, value=sign)
                cells += 1

        # The names have to stay on screen while the eye runs right along thirty columns:
        # that is the one thing the paper table did for free and a spreadsheet does not.
        worksheet.freeze_panes = "B2"
        worksheet.column_dimensions["A"].width = 28
        for column in range(2, len(problems) + 2):
            worksheet.column_dimensions[get_column_letter(column)].width = 5

    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)
    return ExportCounts(sheets=len(sheets), students=len(students), cells=cells)


def _tab_title(number: str) -> str:
    """A worksheet name Excel will accept, built from the листок number.

    Excel refuses ``[]:*?/\\`` in a tab name and truncates past 31 characters.  Sheet
    numbers here are short strings like ``1`` or ``2д`` and hit neither limit, but the
    replacement is done rather than assumed: a number the school invents next year with a
    slash in it would otherwise make ``workbook.save`` raise on a Tuesday evening.
    """
    title = str(number)
    for forbidden in "[]:*?/\\":
        title = title.replace(forbidden, "-")
    return title[:31] or "?"


def out_name(probe: bool = False) -> str:
    """``konduit-2026-09-02.xlsx``, dated in UTC like every other timestamp here."""
    stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    return "konduit-proba-%s.xlsx" % stamp if probe else "konduit-%s.xlsx" % stamp


# ------------------------------------------------------------------------------ the CLI


def command_export(db_path: Optional[Path], out: Optional[Path]) -> int:
    connection = open_read_only(db_path)
    try:
        destination = out or (EXPORT_DIR / out_name())
        counts = export(connection, destination)
    finally:
        connection.close()
    print(counts)
    print("файл: %s" % destination)
    print("в git он не поедет: %s под .gitignore, в нём фамилии 56 детей" % EXPORT_DIR)
    return 0


def command_probe(out: Optional[Path]) -> int:
    """Export a database built here, so the tool can be checked without the live one."""
    with tempfile.TemporaryDirectory(prefix="spetsmat-export-proba-") as directory:
        probe = build_probe_database(Path(directory))
        connection = open_read_only(probe)
        try:
            destination = out or (EXPORT_DIR / out_name(probe=True))
            counts = export(connection, destination)
        finally:
            connection.close()
    print("ПРОБА: база собрана здесь из миграций и засева, живая база НЕ читалась")
    print(counts)
    print("файл: %s" % destination)
    return 0


def main(argv=None) -> int:
    """Returns an exit code, ALWAYS -- a tool that returns None cannot fail physically."""
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--proba", action="store_true",
                        help="выгрузить пробную базу (миграции + засев), не живую")
    parser.add_argument("--baza", type=Path, default=None,
                        help="путь к базе; по умолчанию config.DB_PATH")
    parser.add_argument("--kuda", type=Path, default=None,
                        help="куда положить книгу; по умолчанию data/konduit-<дата>.xlsx")
    arguments = parser.parse_args(argv)

    try:
        if arguments.proba:
            return command_probe(arguments.kuda)
        return command_export(arguments.baza, arguments.kuda)
    except DatabaseMissing as error:
        print(error, file=sys.stderr)
        return 3


if __name__ == "__main__":
    sys.exit(main())
