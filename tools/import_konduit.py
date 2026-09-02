"""Импорт прошлогоднего кондуита в схему бота и сверка проекций.

Проверяем три вещи:
  1. счётчик сдавших в строке 2 каждого листка;
  2. долги из листа «долги»;
  3. гробарий.
Если все три сходятся — модель журнала описывает реальные данные.
"""
import sqlite3
import sys
import openpyxl

SRC = "/sessions/funny-eager-bell/mnt/uploads/Кондуит 8КЛ.xlsx"
SHEETS = ["1", "2", "3", "4", "6", "7", "8", "9", "10", "11",
          "12", "13", "14", "15", "1д", "2д", "3д", "4д"]


def label(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def kind(lbl, status):
    """Тип задачи. Обязательность надёжнее брать из строки статуса:
    в старых листках там '°', в новых '●'. Метка её дублирует не всегда."""
    if status in ("°", "●"):
        return "обязательная"
    if "**" in lbl:
        return "двойная"
    if "*" in lbl:
        return "звезда"
    if "°" in lbl:
        return "обязательная"
    return "обычная"


def layout(ws):
    """Раскладка колонок менялась в течение года — ищем её по заголовкам.

    Возвращает (строка заголовков, колонка фамилии, колонка имени,
    первая колонка задач)."""
    for row in (1, 2, 3, 4):
        headers = {}
        for col in range(1, 12):
            v = ws.cell(row, col).value
            if isinstance(v, str):
                headers[v.strip().lower()] = col
        if "фамилия" in headers and "имя" in headers:
            service = [c for k, c in headers.items()
                       if k in ("фамилия", "имя", "закрыт", "принимающий")]
            return row, headers["фамилия"], headers["имя"], max(service) + 1
    raise ValueError(f"не нашёл заголовков на листе {ws.title!r}")


def value(cell):
    if cell is None:
        return None
    if isinstance(cell, (int, float)) and float(cell) == 1.0:
        return "сдано"
    if isinstance(cell, str) and cell.strip().lower() == "x":
        return "снято"
    return None


def build(db, wb):
    db.executescript("""
        create table students (id integer primary key, фамилия text, имя text);
        create table sheets (id integer primary key, номер text, порядок integer);
        create table problems (id integer primary key, sheet_id int, метка text, тип text);
        create table marks (id integer primary key, student_id int, problem_id int,
                            состояние text);
    """)
    students = {}
    for order, name in enumerate(SHEETS):
        ws = wb[name]
        head_row, col_fam, col_nam, first_col = layout(ws)

        cur = db.execute("insert into sheets (номер, порядок) values (?,?)", (name, order))
        sheet_id = cur.lastrowid

        problems = {}
        for col in range(first_col, ws.max_column + 1):
            lbl = label(ws.cell(1, col).value)
            if not lbl:
                continue
            status = ws.cell(head_row, col).value if head_row > 1 else None
            c = db.execute(
                "insert into problems (sheet_id, метка, тип) values (?,?,?)",
                (sheet_id, lbl, kind(lbl, status)))
            problems[col] = c.lastrowid

        for row in range(head_row + 1, ws.max_row + 1):
            fam = ws.cell(row, col_fam).value
            nam = ws.cell(row, col_nam).value
            if not fam or isinstance(fam, (int, float)):
                continue
            fam = str(fam).strip().lstrip("~ ")
            if not fam:
                continue
            key = (fam, str(nam or "").strip())
            if key not in students:
                c = db.execute("insert into students (фамилия, имя) values (?,?)", key)
                students[key] = c.lastrowid
            sid = students[key]
            for col, pid in problems.items():
                v = value(ws.cell(row, col).value)
                if v:
                    db.execute(
                        "insert into marks (student_id, problem_id, состояние) values (?,?,?)",
                        (sid, pid, v))
    db.commit()
    return students


def check_counters(db, wb):
    """Строка 2 листка против числа сдавших в журнале."""
    bad = []
    for name in SHEETS:
        ws = wb[name]
        head_row, _, _, first_col = layout(ws)
        if head_row == 1:
            continue
        for col in range(first_col, ws.max_column + 1):
            lbl = label(ws.cell(1, col).value)
            expected = ws.cell(2, col).value
            if not lbl or not isinstance(expected, (int, float)):
                continue
            got = db.execute("""
                select count(*) from marks m
                  join problems p on p.id = m.problem_id
                  join sheets s on s.id = p.sheet_id
                 where s.номер = ? and p.метка = ? and m.состояние = 'сдано'
            """, (name, lbl)).fetchone()[0]
            if got != int(expected):
                bad.append((name, lbl, int(expected), got))
    return bad


def check_debts(db, wb):
    """Лист «долги» против несданных обязательных задач."""
    ws = wb["долги"]
    cols = {}
    for col in range(5, ws.max_column + 1):
        lbl = label(ws.cell(1, col).value)
        if lbl:
            cols[col] = lbl
    bad = []
    for row in range(2, ws.max_row + 1):
        fam = ws.cell(row, 1).value
        if not fam:
            continue
        fam = str(fam).strip().lstrip("~ ")
        for col, sheet_no in cols.items():
            cell = ws.cell(row, col).value
            if cell is None:
                continue
            expected = 0 if cell == "✓" else (int(cell) if isinstance(cell, (int, float)) else None)
            if expected is None:
                continue
            got = db.execute("""
                select count(*) from problems p
                  join sheets s on s.id = p.sheet_id
                 where s.номер = ? and p.тип = 'обязательная'
                   and not exists (
                       select 1 from marks m
                        join students st on st.id = m.student_id
                        where m.problem_id = p.id and st.фамилия = ?)
            """, (sheet_no, fam)).fetchone()[0]
            if got != expected:
                bad.append((fam, sheet_no, expected, got))
    return bad


def check_graveyard(db, wb):
    """Гробарий: сколько человек взяло задачу и какой знак стоит."""
    ws = wb["гробарий"]
    rows = []
    for row in range(2, ws.max_row + 1):
        sheet_no = ws.cell(row, 1).value
        lbl = label(ws.cell(row, 2).value)
        sign = ws.cell(row, 3).value
        if not sheet_no or not lbl:
            continue
        sheet_no = str(sheet_no).strip("[] ")
        named = sum(1 for c in range(4, ws.max_column + 1) if ws.cell(row, c).value)
        got = db.execute("""
            select count(*) from marks m
              join problems p on p.id = m.problem_id
              join sheets s on s.id = p.sheet_id
             where s.номер = ? and p.метка = ? and m.состояние = 'сдано'
        """, (sheet_no, lbl)).fetchone()[0]
        rows.append((sheet_no, lbl, sign, named, got))
    return rows


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    db = sqlite3.connect(":memory:")
    students = build(db, wb)

    n_marks = db.execute("select count(*) from marks").fetchone()[0]
    n_probs = db.execute("select count(*) from problems").fetchone()[0]
    print(f"учеников {len(students)} · задач {n_probs} · отметок {n_marks}\n")

    bad = check_counters(db, wb)
    print(f"[1] счётчики листков: расхождений {len(bad)}")
    for b in bad[:10]:
        print(f"    листок {b[0]} задача {b[1]}: в таблице {b[2]}, в журнале {b[3]}")

    bad = check_debts(db, wb)
    print(f"\n[2] долги: расхождений {len(bad)}")
    for b in bad[:10]:
        print(f"    {b[0]} листок {b[1]}: в таблице {b[2]}, в журнале {b[3]}")

    print("\n[3] гробарий: листок · задача · знак · названо имён · сдавших")
    for r in check_graveyard(db, wb):
        print(f"    {r[0]:>4} {r[1]:>6}  {r[2]}  имён {r[3]}  сдавших {r[4]}")


if __name__ == "__main__":
    sys.exit(main())
