"""The export: it agrees with the screen, it uses the conduit's signs, and it cannot write.

The three things this suite is actually about map one-to-one onto the three ways an export
of a journal goes wrong: it disagrees with the projection everyone else reads, it invents
an alphabet nobody can read back, or it opens the live database for writing during a
lesson.  Everything else here is arithmetic around those three.
"""

from __future__ import annotations

import sqlite3

import pytest
from openpyxl import load_workbook

from core.models import CellState
from core.services.progress import ProgressService
from infra.repositories import SqliteCatalogue, SqliteMarkJournal
from tools.export_xlsx import (
    SIGN,
    DatabaseMissing,
    ExportCounts,
    _tab_title,
    export,
    main,
    open_read_only,
    out_name,
)


# ------------------------------------------------------------------ agrees with the screen


def test_kazhdaya_kletka_sovpadaet_s_proekciej(connection, raznocvetnyy_mir, tmp_path):
    """Every cell of the workbook says what ``ProgressService`` says, cell by cell.

    This is the assertion the whole export exists to keep.  It is written as a sweep over
    the entire rectangle rather than over three interesting cells, and it reports its own
    coverage, because "the export matches the projection" checked on one cell and checked
    on all of them read identically in a green log.
    """
    out = tmp_path / "konduit.xlsx"
    counts = export(connection, out)

    catalogue = SqliteCatalogue(connection)
    progress = ProgressService(SqliteMarkJournal(connection), catalogue)
    students = catalogue.students()
    workbook = load_workbook(out)

    sverennyh = 0
    for sheet in catalogue.sheets():
        problems = catalogue.problems_of_sheet(sheet.id)
        states = progress.states_for_many([s.id for s in students], [p.id for p in problems])
        worksheet = workbook[_tab_title(sheet.number)]
        for row, student in enumerate(students, start=2):
            for column, problem in enumerate(problems, start=2):
                v_knige = worksheet.cell(row=row, column=column).value or ""
                assert v_knige == SIGN[states[(student.id, problem.id)]], (
                    "лист %s, %s %s, задача %s"
                    % (sheet.number, student.surname, student.name, problem.label)
                )
                sverennyh += 1

    assert sverennyh == counts.cells, "сверено %d клеток из %d" % (sverennyh, counts.cells)
    assert sverennyh > 0


def test_tri_sostoyaniya_dayut_tri_raznyh_znaka(connection, raznocvetnyy_mir, tmp_path):
    """SOLVED, RETRACTED and EMPTY reach the workbook as ``1``, ``x`` and a blank.

    Named separately from the sweep above: the sweep compares the export against ``SIGN``,
    so it would stay green if ``SIGN`` itself were rewritten to put ``1`` under every
    state.  This test names the three characters out loud.
    """
    out = tmp_path / "konduit.xlsx"
    export(connection, out)

    worksheet = load_workbook(out)["1"]
    znaki = {
        worksheet.cell(row=row, column=column).value
        for row in range(2, 5)
        for column in range(2, 6)
    }
    assert "1" in znaki
    assert "x" in znaki
    assert None in znaki, "пустая клетка обязана остаться пустой, а не стать нулём"


def test_pustaya_posle_erratum_pusta_a_ne_plyus(connection, raznocvetnyy_mir, tmp_path):
    """A cell struck out by ``erratum`` is blank, not the plus that stood there before.

    The projection never looks past the last event, and the export must not either: a
    workbook that showed the pre-erratum plus would be the one place in the system where a
    wrong button survives being corrected.
    """
    out = tmp_path / "konduit.xlsx"
    export(connection, out)

    worksheet = load_workbook(out)["1"]
    # Third student, third problem: given, then struck out -- see ``raznocvetnyy_mir``.
    assert worksheet.cell(row=4, column=4).value is None


# --------------------------------------------------------------------------- the shape


def test_odin_list_na_listok_ucheniki_vniz_zadachi_vpravo(connection, malenkiy_mir, tmp_path):
    out = tmp_path / "konduit.xlsx"
    counts = export(connection, out)
    workbook = load_workbook(out)

    catalogue = SqliteCatalogue(connection)
    sheets = catalogue.sheets()
    assert workbook.sheetnames == [_tab_title(s.number) for s in sheets]

    worksheet = workbook[_tab_title(sheets[0].number)]
    problems = catalogue.problems_of_sheet(sheets[0].id)
    assert worksheet.cell(row=1, column=1).value == "Ученик"
    assert [worksheet.cell(row=1, column=c).value for c in range(2, len(problems) + 2)] == [
        p.label for p in problems
    ]
    students = catalogue.students()
    assert [worksheet.cell(row=r, column=1).value for r in range(2, len(students) + 2)] == [
        "%s %s" % (s.surname, s.name) for s in students
    ]
    assert counts == ExportCounts(
        sheets=len(sheets), students=len(students),
        cells=len(students) * sum(len(catalogue.problems_of_sheet(s.id)) for s in sheets),
    )


def test_listok_bez_zadach_vsyo_ravno_poluchaet_list(connection, tmp_path):
    """A листок with no problems is still a worksheet.

    A sheet missing from the workbook is indistinguishable from a sheet that never
    existed, and "the листок we did not do yet" is a fact the owner keeps in that book.
    """
    connection.execute(
        "insert into sheets (number, title, issued_at, ord) values ('7', 'пустой', '2026-09-01', 7)"
    )
    connection.commit()
    export(connection, tmp_path / "konduit.xlsx")
    assert "7" in load_workbook(tmp_path / "konduit.xlsx").sheetnames


@pytest.mark.parametrize(
    "number, expected",
    [("1", "1"), ("2д", "2д"), ("a/b", "a-b"), ("a:b*c", "a-b-c"), ("я" * 40, "я" * 31)],
)
def test_imya_lista_prohodit_ogranicheniya_excel(number, expected):
    assert _tab_title(number) == expected


def test_imya_fajla_nesyot_datu():
    assert out_name().startswith("konduit-")
    assert out_name().endswith(".xlsx")
    assert "proba" in out_name(probe=True)


# ------------------------------------------------------------------------ cannot write


def test_soedinenie_tolko_na_chtenie(zakrytaya_baza):
    """A write through the export's own connection raises instead of landing in the file.

    The property the задание states as "export is a READ, safe to run while teachers are
    marking", tested against the connection rather than against a comment about it.

    The MESSAGE is not asserted, only the refusal and the untouched bytes: SQLite says
    "attempt to write a readonly database" for a plain file and "unable to open database
    file" for one in WAL mode, where the write would first have to create the ``-wal``
    beside it.  Both are refusals; pinning either one would make the test a statement
    about the journal mode instead of about the write.
    """
    do = zakrytaya_baza.read_bytes()
    connection = open_read_only(zakrytaya_baza)
    try:
        with pytest.raises(sqlite3.OperationalError):
            connection.execute(
                "insert into sheets (number, title, issued_at, ord) "
                "values ('99', 'нет', '2026-01-01', 99)"
            )
    finally:
        connection.close()
    assert zakrytaya_baza.read_bytes() == do


def test_eksport_ne_menyaet_bajty_bazy(connection, raznocvetnyy_mir, zakrytaya_baza, tmp_path):
    """The database file is byte-identical after an export.

    Not "the data is the same": the FILE.  An export that rewrote the database in a way
    that happened to preserve the rows would still be an export that wrote to the file
    while a lesson was running.
    """
    do = zakrytaya_baza.read_bytes()
    otkrytoe = open_read_only(zakrytaya_baza)
    try:
        export(otkrytoe, tmp_path / "konduit.xlsx")
    finally:
        otkrytoe.close()
    assert zakrytaya_baza.read_bytes() == do


def test_bez_bazy_govorit_slovami_a_ne_tracebackom(tmp_path):
    with pytest.raises(DatabaseMissing) as error:
        open_read_only(tmp_path / "нет-такой.db")
    assert "нет базы" in str(error.value)
    assert "--proba" in str(error.value), "сообщение обязано называть, чем это проверить"


# -------------------------------------------------------------------------------- CLI


def test_proba_vygruzhaet_ves_zasev(tmp_path, capsys):
    """``--proba`` is the criterion's command: rc=0 and the three measured numbers.

    Eighteen листков and fifty-six students are what ``seed/`` contains, so this is also
    the test that the export walks the whole catalogue and not the first page of it.
    """
    out = tmp_path / "proba.xlsx"
    assert main(["--proba", "--kuda", str(out)]) == 0

    vyvod = capsys.readouterr().out
    assert "листов 18, учеников 56, клеток 30464" in vyvod
    assert str(out) in vyvod
    assert out.exists()
    assert len(load_workbook(out).sheetnames) == 18


def test_bez_bazy_cli_vozvrashchaet_3(tmp_path, capsys):
    assert main(["--baza", str(tmp_path / "нет.db")]) == 3
    assert "нет базы" in capsys.readouterr().err


def test_cli_vsegda_vozvrashchaet_kod(tmp_path, raznocvetnyy_mir, zakrytaya_baza):
    """A tool that returns ``None`` cannot fail physically; this one returns an int."""
    kod = main(["--baza", str(zakrytaya_baza), "--kuda", str(zakrytaya_baza.parent / "k.xlsx")])
    assert isinstance(kod, int) and kod == 0


def test_znaki_eto_alfavit_konduita():
    """The signs are the importer's registry read the other way, not a local invention."""
    assert SIGN[CellState.SOLVED] == "1"
    assert SIGN[CellState.RETRACTED] == "x"
    assert SIGN[CellState.EMPTY] == ""
    assert set(SIGN) == set(CellState), "у каждого состояния обязан быть свой знак"
